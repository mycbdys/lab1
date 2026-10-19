"""Build the simple-English version of the control-questions document."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "Lab1_Control_Questions_Simple.docx"

BLUE = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x21, 0x25, 0x29)
GREY = RGBColor(0x59, 0x59, 0x59)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    for tag, value in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tag)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = value
            run._r.append(instr)


def style(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

    h = doc.styles["Heading 2"]
    h.font.name = "Calibri"
    h.font.size = Pt(11.5)
    h.font.bold = True
    h.font.color.rgb = BLUE
    h.paragraph_format.space_before = Pt(11)
    h.paragraph_format.space_after = Pt(3)
    h.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("HDU IS - Laboratory Work 1 | Control Questions (simple) | page ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    add_page_number(footer)


def header_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("HDU IS — Laboratory Work 1")
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    r.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("Control Questions — Simple Answers")
    r.font.size = Pt(22)
    r.font.color.rgb = BLUE
    r.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("Secure REST API Development with CI/CD and Real CVE Vulnerability Analysis")
    r.font.size = Pt(9.5)
    r.font.color.rgb = GREY
    r.italic = True

    info = doc.add_paragraph()
    info.paragraph_format.space_after = Pt(8)
    r = info.add_run(
        "李泽涛 / LiZetao  |  HDU 23320332  |  Group 23320113  |  Variant N/A  |  19 October 2026"
    )
    r.font.size = Pt(9.5)
    r.font.color.rgb = DARK


def main() -> None:
    doc = Document()
    style(doc)
    header_block(doc)

    text = (HERE / "Control_Questions_Simple.md").read_text(encoding="utf-8")
    parts = re.split(r"^## ", text, flags=re.MULTILINE)[1:]
    for part in parts:
        lines = part.strip().split("\n")
        title = lines[0].strip()
        match = re.match(r"^(\d+)\.\s*(.*)$", title)
        if match:
            title = f"Q{match.group(1)}. {match.group(2)}"
        doc.add_heading(title, level=2)
        body = "\n".join(lines[1:]).strip()
        for paragraph_text in re.split(r"\n\s*\n", body):
            cleaned = " ".join(paragraph_text.split())
            cleaned = cleaned.replace("`", "").replace("**", "")
            if cleaned:
                doc.add_paragraph(cleaned)

    doc.core_properties.title = "Work 1 - Control Questions (simple answers)"
    doc.core_properties.author = "李泽涛 (LiZetao)"
    doc.core_properties.subject = "HDU IS Laboratory Work 1"
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
