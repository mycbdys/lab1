"""Build a concise version of the Work 1 report (about 5-6 pages)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "Lab1_Report.docx"

BLUE = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x21, 0x25, 0x29)
GREY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xB0, 0x2A, 0x37)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    for tag, value in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tag)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = value
            run._r.append(instr)


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.left_indent = Cm(0.2)
    p.paragraph_format.keep_together = True
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "F4F6F8")
    p_pr.append(shd)
    lines = text.strip("\n").split("\n")
    for index, line in enumerate(lines):
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8)
        run.font.color.rgb = DARK
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        if index != len(lines) - 1:
            run.add_break()


def add_table(doc: Document, headers, rows, widths=None, size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(header)
        run.bold = True
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tc_pr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "1F4E79")
        tc_pr.append(shd)
    for row_data in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row_data):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(str(value))
            run.font.size = Pt(size)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
    if widths:
        for index, width in enumerate(widths):
            table.columns[index].width = Cm(width)
            for cell in table.columns[index].cells:
                cell.width = Cm(width)
    return table


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREY


def style(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.1)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.10
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

    for name, size, before, after in (
        ("Heading 1", 14, 14, 5),
        ("Heading 2", 11.5, 10, 3),
    ):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = BLUE
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("HDU IS - Laboratory Work 1 | page ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    add_page_number(footer)


def title_page(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("HANGZHOU DIANZI UNIVERSITY")
    r.font.size = Pt(12)
    r.font.color.rgb = GREY
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Information Security Laboratory")
    r.font.size = Pt(12)
    r.font.color.rgb = GREY

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("LABORATORY WORK No. 1")
    r.font.size = Pt(13)
    r.font.color.rgb = RED
    r.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("Secure REST API Development with CI/CD\nand Real CVE Vulnerability Analysis")
    r.font.size = Pt(23)
    r.font.color.rgb = BLUE
    r.bold = True

    doc.add_paragraph()
    doc.add_paragraph()

    info = [
        ("Variant", "N/A (not assigned)"),
        ("Student", "李泽涛 / LiZetao"),
        ("HDU number", "23320332"),
        ("Group", "23320113"),
        ("Repository", "https://github.com/mycbdys/lab1"),
        ("Submission date", "19 October 2026"),
    ]
    table = doc.add_table(rows=0, cols=2)
    for label, value in info:
        cells = table.add_row().cells
        r = cells[0].paragraphs[0].add_run(label)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = BLUE
        cells[1].text = value
        for p in cells[1].paragraphs:
            for run in p.runs:
                run.font.size = Pt(10.5)
    for row in table.rows:
        row.cells[0].width = Cm(3.8)
        row.cells[1].width = Cm(10.5)

    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Python 3.13 / Flask 3.1  |  SQLite  |  pytest  |  Bandit  |  pip-audit  |  Node.js 20 / EJS")
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREY

    doc.add_page_break()


def main() -> None:
    doc = Document()
    style(doc)
    title_page(doc)

    doc.add_heading("1. Summary", level=1)
    doc.add_paragraph(
        "This work implements a small secure REST API in Python/Flask, runs automated SAST and SCA "
        "checks in GitHub Actions, and analyses a real vulnerability: CVE-2022-29078, a remote-code-"
        "execution flaw in the EJS template engine. The API has three endpoints, bcrypt password "
        "storage, JWT authentication, parameterised SQL and output encoding. The CVE was reproduced "
        "in a controlled environment, its root cause was identified in the EJS source, and the fix "
        "(EJS 3.1.7 plus safer application code) was verified."
    )
    doc.add_paragraph(
        "Repository: https://github.com/mycbdys/lab1  |  CI: "
        "https://github.com/mycbdys/lab1/actions/runs/39195640683"
    )

    doc.add_heading("2. API and security controls", level=1)
    add_table(
        doc,
        ["Endpoint", "Auth", "Purpose"],
        [
            ("POST /auth/login", "No", "Verify credentials, return a 15-minute JWT"),
            ("GET /api/data", "Bearer JWT", "Return only the caller's posts, HTML-escaped"),
            ("POST /api/posts", "Bearer JWT", "Create a post owned by the authenticated user"),
        ],
        widths=[4.2, 3.0, 9.6],
    )
    doc.add_paragraph(
        "SQL injection is prevented with bound parameters for every query; user input is never "
        "concatenated into SQL. Stored XSS is prevented by HTML-escaping all text on output and by a "
        "strict Content-Security-Policy. Passwords are hashed with bcrypt (cost 12, unique salt); "
        "login errors are identical for unknown users and wrong passwords, and five failed attempts "
        "per minute return 429. JWTs are HS256 with a pinned algorithm, issuer, audience and required "
        "exp/iat/sub claims, so tampered, expired and alg=none tokens are rejected. Every data query "
        "is scoped by the user id from the verified token, and responses carry nosniff, DENY frame "
        "options, no-referrer and no-store headers."
    )

    doc.add_heading("3. CI/CD and results", level=1)
    add_table(
        doc,
        ["Job", "Tool", "Gate"],
        [
            ("Unit tests", "pytest 9.1.1", "21 tests must pass"),
            ("SAST", "Bandit 1.9.4", "No medium/high findings"),
            ("SCA", "pip-audit 2.10.1", "No known vulnerable dependency"),
        ],
        widths=[3.4, 5.0, 8.4],
    )
    doc.add_paragraph(
        "The workflow runs on every push and pull request. Results: 21/21 tests passed, Bandit "
        "reported 0 issues in 303 lines of code, and pip-audit found 0 known vulnerabilities in 44 "
        "packages. The same pipeline passed on GitHub's clean runner, so the result is not only local."
    )

    doc.add_heading("4. CVE-2022-29078", level=1)
    add_table(
        doc,
        ["Field", "Value"],
        [
            ("CVE", "CVE-2022-29078 (GHSA-phwq-j96m-2c2q)"),
            ("Component", "EJS template engine for Node.js"),
            ("Affected / fixed", "EJS < 3.1.7 / EJS 3.1.7"),
            ("Class / severity", "CWE-74 injection leading to RCE / CVSS 9.8 Critical"),
        ],
        widths=[4.0, 12.8],
    )
    doc.add_paragraph(
        "The vulnerable app forwards the whole query object into res.render(). EJS merges "
        "data.settings['view options'] into its compiler options, and outputFunctionName is "
        "concatenated directly into the generated JavaScript. A crafted request therefore executes a "
        "command while the page still returns 200."
    )
    add_code(
        doc,
        """GET /?settings[view options][outputFunctionName]=
x;process.mainModule.require("child_process").execSync("id > pwned.txt");s""",
    )
    doc.add_paragraph(
        "Observed in Docker (GitHub Actions): HTTP 200 with the normal page, and pwned.txt inside "
        "the vulnerable container contained uid=0(root). The same request against the fixed app "
        "(EJS 3.1.7, scalar-only render data) returned the normal page and created no file; the "
        "library also rejected the option as 'not a valid JS identifier'."
    )
    doc.add_heading("Root cause and fix", level=2)
    doc.add_paragraph(
        "Three conditions combine: (1) the application passes attacker-controlled data as render "
        "options; (2) EJS copies settings['view options'] into internal compiler options without an "
        "allow-list; (3) outputFunctionName is inserted into generated source, so the payload becomes "
        "executable statements. The fix is applied at both layers: upgrade to EJS 3.1.7, which "
        "validates the option as a JavaScript identifier, and stop passing the query object, passing "
        "only a 64-character scalar name. Neither fix relies on a blocklist of characters."
    )

    doc.add_heading("5. Evidence (screenshots)", level=1)
    doc.add_paragraph(
        "The screenshots below were captured from the public GitHub repository and Actions pages; "
        "they are not mock-ups. The first shows the three CI jobs passing, and the second shows the "
        "Docker run steps plus the files produced inside the containers."
    )
    doc.add_picture(str(HERE / "images" / "screenshots" / "combined-ci.png"), width=Cm(16.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(
        doc,
        "Figure 1. GitHub Actions: Unit tests, SAST (Bandit) and SCA (pip-audit) all pass; the Docker "
        "job builds the vulnerable stack, reproduces the CVE, starts the fixed stack and verifies it.",
    )
    doc.add_picture(str(HERE / "images" / "screenshots" / "combined-cve.png"), width=Cm(16.2))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(
        doc,
        "Figure 2. Docker evidence: pwned.txt from the vulnerable container contains uid=0(root); "
        "the fixed container runs EJS 3.1.7.",
    )

    doc.add_heading("6. Testing and submission", level=1)
    doc.add_paragraph(
        "The pytest suite covers login, SQL injection attempts, bcrypt storage, rate limiting, "
        "missing/tampered/expired/alg=none tokens, per-user isolation, XSS escaping, SQL "
        "metacharacters stored literally, and security headers. Manual cURL evidence is in "
        "evidence/api-demo.txt; CVE evidence is in cve/CVE-2022-29078/evidence/ and "
        "docs/evidence/docker-ci/."
    )
    doc.add_paragraph(
        "Deliverables: public repository https://github.com/mycbdys/lab1; pipeline "
        "https://github.com/mycbdys/lab1/actions/runs/39195640683; control questions in the "
        "separate document Lab1_Control_Questions.docx. The fixes are under "
        "cve/CVE-2022-29078/fixed/ with patches/fix.diff."
    )
    doc.add_paragraph(
        "Limitation: Docker Desktop is not installed on the local Windows host (it needs "
        "administrator rights, WSL2 and a reboot). The same compose files were executed in real "
        "Docker on the GitHub Actions runner (run 39195350651): the vulnerable container ran id as "
        "root and the fixed container blocked the attack. The local run used Node.js 20 with the "
        "same package versions."
    )

    doc.core_properties.title = "Secure REST API Development with CI/CD and Real CVE Vulnerability Analysis"
    doc.core_properties.author = "李泽涛 (LiZetao)"
    doc.core_properties.subject = "HDU IS Laboratory Work 1 - concise report"
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
