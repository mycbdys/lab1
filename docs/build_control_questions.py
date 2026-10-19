"""Build the standalone control-questions document."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
OUT = HERE / "Lab1_Control_Questions.docx"

BLUE = RGBColor(0x1F, 0x4E, 0x79)
DARK = RGBColor(0x21, 0x25, 0x29)
GREY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xB0, 0x2A, 0x37)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    for tag, value in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), tag)
            run._r.append(fld)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = value
            run._r.append(instr)


def style(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DARK
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.14
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

    for name, size, before, after in (
        ("Heading 1", 15, 18, 8),
        ("Heading 2", 12, 14, 5),
    ):
        h = doc.styles[name]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = BLUE
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("HDU IS - Laboratory Work 1 | Control Questions | page ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    add_page_number(footer)


def title_page(doc: Document) -> None:
    for _ in range(3):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("HANGZHOU DIANZI UNIVERSITY")
    run.font.size = Pt(12)
    run.font.color.rgb = GREY
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Information Security Laboratory")
    run.font.size = Pt(12)
    run.font.color.rgb = GREY

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("LABORATORY WORK No. 1")
    run.font.size = Pt(13)
    run.font.color.rgb = RED
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    run = p.add_run("Control Questions")
    run.font.size = Pt(30)
    run.font.color.rgb = BLUE
    run.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Secure REST API Development with CI/CD and Real CVE Vulnerability Analysis")
    run.font.size = Pt(12)
    run.font.color.rgb = GREY
    run.italic = True

    doc.add_paragraph()
    doc.add_paragraph()

    info = [
        ("Variant", "N/A (not assigned)"),
        ("Student", "李泽涛 / LiZetao"),
        ("HDU number", "23320332"),
        ("Group", "23320113"),
        ("Submission date", "19 October 2026"),
    ]
    table = doc.add_table(rows=0, cols=2)
    for label, value in info:
        cells = table.add_row().cells
        run = cells[0].paragraphs[0].add_run(label)
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = BLUE
        cells[1].text = value
        for paragraph in cells[1].paragraphs:
            for r in paragraph.runs:
                r.font.size = Pt(10.5)
    for row in table.rows:
        row.cells[0].width = Cm(4.0)
        row.cells[1].width = Cm(9.5)

    doc.add_page_break()


def main() -> None:
    doc = Document()
    style(doc)
    title_page(doc)

    doc.add_heading("Answers to the eight control questions", level=1)
    doc.add_paragraph(
        "The answers below refer to the implementation, the CI/CD pipeline and the CVE case study "
        "delivered in the main report (Lab1_Report.docx). The main report itself does not duplicate "
        "these answers."
    )

    text = (HERE / "Control_Questions.md").read_text(encoding="utf-8")
    parts = re.split(r"^## ", text, flags=re.MULTILINE)[1:]
    for part in parts:
        lines = part.strip().split("\n")
        title = lines[0].strip()
        match = re.match(r"^(\d+)\.\s*(.*)$", title)
        if match:
            title = f"Q{match.group(1)}. {match.group(2)}"
        doc.add_heading(title, level=2)
        body = "\n".join(lines[1:]).strip()
        for paragraph_text in re.split(r"\n\s*\n", body):
            cleaned = " ".join(paragraph_text.split())
            cleaned = cleaned.replace("`", "").replace("**", "")
            if cleaned:
                doc.add_paragraph(cleaned)

    doc.core_properties.title = "Work 1 - Control Questions"
    doc.core_properties.author = "李泽涛 (LiZetao)"
    doc.core_properties.subject = "HDU IS Laboratory Work 1"
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
